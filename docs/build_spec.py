import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
BLACK = RGBColor(0, 0, 0)


def element(tag, **attrs):
    item = OxmlElement(tag)
    for key, value in attrs.items():
        item.set(qn(key), str(value))
    return item


def field(paragraph, instruction):
    run = paragraph.add_run()
    run._r.append(element("w:fldChar", **{"w:fldCharType": "begin"}))
    text = element("w:instrText", **{"xml:space": "preserve"})
    text.text = f" {instruction} "
    run._r.append(text)
    run._r.append(element("w:fldChar", **{"w:fldCharType": "end"}))


def hyperlink(paragraph, label, url=None, anchor=None):
    link = element("w:hyperlink")
    if url:
        rel = paragraph.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )
        link.set(qn("r:id"), rel)
    else:
        link.set(qn("w:anchor"), anchor)
    run = element("w:r")
    properties = element("w:rPr")
    properties.append(element("w:color", **{"w:val": "175580" if url else "000000"}))
    properties.append(element("w:u", **{"w:val": "single" if url else "none"}))
    run.append(properties)
    text = element("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, text):
    text = text.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    pattern = r"(\[[^\]]+\]\([^\s)]+\)|\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))"
    for part in re.split(pattern, text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^\s)]+)\)", part)
        if link:
            hyperlink(paragraph, link[1], url=link[2])
            continue
        run = paragraph.add_run(part.strip("*") if part.startswith("*") else part.strip("`") if part.startswith("`") else part)
        if part.startswith("**"):
            run.bold = True
        elif part.startswith("*"):
            run.italic = True
        elif part.startswith("`"):
            run.font.name = "DejaVu Sans Mono"
            run.font.size = Pt(9)


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    sec.left_margin, sec.right_margin = Cm(1.85), Cm(1.85)
    sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True
    rpr = normal.element.get_or_add_rPr()
    rpr.append(element("w:lang", **{"w:val": "fr-FR"}))
    for name, size, before, after in [
        ("Title", 26, 0, 14),
        ("Subtitle", 13, 0, 9),
        ("Heading 1", 17, 16, 8),
        ("Heading 2", 13, 11, 6),
        ("Heading 3", 11, 9, 4),
        ("Heading 4", 10.5, 7, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.page_break_before = False
        for color in style.element.xpath(".//w:color"):
            for attr in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn(f"w:{attr}"), None)
        for borders in style.element.xpath(".//w:pBdr"):
            borders.getparent().remove(borders)
    for name in ("List Bullet", "List Number", "List Bullet 2", "List Number 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.06
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Cahier des charges  |  ")
    field(footer, "PAGE")
    footer.add_run(" / ")
    field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = BLACK
    doc.core_properties.title = "Cahier des charges de gestion patrimoniale et locative"
    doc.core_properties.subject = "Spécification produit et technique de la surcouche SaaS connectée à Odoo"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "SCI, patrimoine, gestion locative, Odoo, SaaS"
    doc.core_properties.language = "fr-FR"
    return doc


def cells(line):
    return [value.strip() for value in re.split(r"(?<!\\)\|", line.strip().strip("|"))]


def table_widths(rows):
    count = len(rows[0])
    if count == 2:
        return [4.1, 13.2]
    scores = []
    for col in range(count):
        values = [len(re.sub(r"[*`]", "", row[col])) for row in rows if col < len(row)]
        mean = sum(values) / len(values)
        scores.append(max(15, min(90, mean)) ** 0.65)
    total = sum(scores)
    return [17.3 * score / total for score in scores]


def add_table(doc, rows):
    width = len(rows[0])
    if any(len(row) != width for row in rows):
        raise ValueError(f"Inconsistent table width: {rows[0]}")
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = table_widths(rows)
    for col, value in zip(table.columns, widths):
        col.width = Cm(value)
    properties = table._tbl.tblPr
    borders = element("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(element(f"w:{edge}", **{"w:val": "single", "w:sz": "4", "w:color": "D9D9D9"}))
    properties.append(borders)
    margins = element("w:tblCellMar")
    for edge, value in (("top", 75), ("bottom", 75), ("left", 100), ("right", 100)):
        margins.append(element(f"w:{edge}", **{"w:w": value, "w:type": "dxa"}))
    properties.append(margins)
    for r, source in enumerate(rows):
        row = table.rows[r]
        row._tr.get_or_add_trPr().append(element("w:cantSplit"))
        if r == 0:
            row._tr.get_or_add_trPr().append(element("w:tblHeader"))
        for c, text in enumerate(source):
            cell = row.cells[c]
            cell.width = Cm(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill = "E7EDF2" if r == 0 else "FFFFFF" if r % 2 else "F6F8FA"
            cell._tc.get_or_add_tcPr().append(element("w:shd", **{"w:fill": fill}))
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.04
            paragraph.paragraph_format.keep_with_next = r == 0
            if r > 0 and len(text) < 16 and re.fullmatch(r"[\w /+−–-]+", text):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if r == 0:
                    run.bold = True
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = Pt(4)
    gap.add_run().font.size = Pt(4)


def add_bookmark(paragraph, index):
    paragraph._p.insert(0, element("w:bookmarkStart", **{"w:id": index, "w:name": f"section_{index}"}))
    paragraph._p.append(element("w:bookmarkEnd", **{"w:id": index}))


def build(source, output):
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = setup_document()
    major = [re.sub(r"^##\s+", "", line) for line in lines if re.match(r"^##\s+", line)]
    i, section_index, title_seen = 0, 0, False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue
        if line == "<!-- toc -->":
            doc.add_paragraph("Sommaire", "Heading 1")
            for index, heading in enumerate(major, start=1):
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(3)
                hyperlink(paragraph, heading, anchor=f"section_{index}")
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            for entry in block:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1
                run = paragraph.add_run(entry)
                run.font.name = "DejaVu Sans Mono"
                run.font.size = Pt(8.5)
            i += 1
            continue
        heading = re.match(r"^(#{1,5})\s+(.*)$", line)
        if heading:
            depth, text = len(heading[1]), heading[2]
            if depth == 1 and not title_seen:
                paragraph = doc.add_paragraph(style="Title")
                title_seen = True
            else:
                paragraph = doc.add_paragraph(style=f"Heading {min(max(depth - 1, 1), 4)}")
            inline(paragraph, text)
            if depth == 2:
                section_index += 1
                add_bookmark(paragraph, section_index)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?[ :|-]+\|[ :|-]*$", lines[i + 1]):
            rows = [cells(line)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(cells(lines[i]))
                i += 1
            add_table(doc, rows)
            continue
        bullet = re.match(r"^(\s*)([-*]|\d+[.)])\s+(.*)$", lines[i])
        if bullet:
            level = " 2" if len(bullet[1]) > 1 else ""
            is_numbered = bullet[2][0].isdigit()
            if is_numbered:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.65 if not level else 1.3)
                paragraph.paragraph_format.first_line_indent = Cm(-0.5)
                inline(paragraph, f"{bullet[2]}  {bullet[3]}")
            else:
                paragraph = doc.add_paragraph(style=f"List Bullet{level}")
                inline(paragraph, bullet[3])
            i += 1
            continue
        if re.fullmatch(r"[-_*]{3,}", line):
            i += 1
            continue
        paragraph = doc.add_paragraph()
        paragraph_lines = [line.lstrip("> ") if line.startswith(">") else line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"\s*(#|\||```|<!--|[-*] |\d+[.)] )", lines[i]):
            paragraph_lines.append(lines[i].strip())
            i += 1
        inline(paragraph, " ".join(paragraph_lines))
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Created {output}")
    print(f"Paragraphs: {len(doc.paragraphs)}; tables: {len(doc.tables)}; major sections: {section_index}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", nargs="?", type=Path, default=ROOT / "cahier-des-charges-gestion-patrimoniale.md")
    parser.add_argument("output", nargs="?", type=Path, default=ROOT / "cahier-des-charges-gestion-patrimoniale.docx")
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
